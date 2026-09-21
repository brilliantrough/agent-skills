"""Deliberately bounded SVG path converter; see SKILL.md before extending its profile."""
from collections import Counter
from copy import deepcopy
from io import BytesIO
from math import atan2, degrees, isclose, prod
import re
from zipfile import ZipFile

import lxml.etree as ET
from PIL import Image, ImageColor
from fontTools.svgLib.path import parse_path
from fontTools.pens.recordingPen import RecordingPen
from fontTools.pens.boundsPen import BoundsPen
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.shapes.autoshape import Shape

from convert import require, svg_size


PAINT = {'fill', 'stroke', 'stroke-width', 'stroke-linecap', 'stroke-linejoin', 'fill-rule'}
META = {'id', 'class', 'label', 'data-name'}


def attrs(e):
    result = dict(e.attrib)
    for entry in e.get('style', '').split(';'):
        if entry.strip():
            key, value = entry.split(':', 1)
            result[key.strip()] = value.strip()
    result.pop('style', None)
    return result


def validate_svg(root, width, height):
    require(max(width, height) <= 56, 'PowerPoint slide size exceeds 56 inches')
    box = [float(v) for v in re.split(r'[ ,]+', root.get('viewBox', '').strip()) if v]
    require(len(box) == 4 and box[:2] == [0, 0] and min(box[2:]) > 0,
            'PPTX requires a positive viewBox starting at 0,0')
    require(isclose(box[2]/box[3], width/height, rel_tol=1e-5), 'Nonuniform SVG/page scaling is not supported')
    supported = {'svg', 'g', 'defs', 'path', 'rect', 'text', 'tspan', 'title', 'desc',
                 'linearGradient', 'stop', 'filter', 'feGaussianBlur'}
    for e in root.iter():
        kind = ET.QName(e).localname
        require(kind in supported, f'Unsupported SVG element: {kind}; flatten geometry first')
        a = attrs(e)
        for key in ('clip-path', 'mask', 'display', 'visibility', 'fill-opacity', 'stroke-opacity'):
            require(key not in a, f'{key} is not supported; do not silently discard it')
        filtered_parent = any(p.get('filter') for p in e.iterancestors())
        if 'transform' in a:
            require(filtered_parent and kind == 'path' and re.fullmatch(r'translate\([\d. ,+-]+\)', a['transform']),
                    'Transforms outside a rasterized shadow are not supported')
        if 'filter' in a:
            require(kind == 'g' and not filtered_parent and len(e) == 1 and ET.QName(e[0]).localname == 'path',
                    'Only a single-path shadow group may be rasterized')
        if 'opacity' in a:
            require('filter' in a and 0 <= float(a['opacity']) <= 1, 'Only shadow-group opacity is supported')
        if kind == 'svg':
            require(e is root, 'Nested SVG viewports are not supported')
            allowed = PAINT | META | {'width', 'height', 'viewBox', 'preserveAspectRatio', 'version', 'space'}
            require(all(ET.QName(k).localname in allowed for k in a), 'Unsupported root SVG attributes')
        if kind in ('path', 'g', 'rect'):
            allowed = PAINT | META | {'d', 'x', 'y', 'width', 'height', 'transform', 'filter', 'opacity'}
            require(all(ET.QName(k).localname in allowed for k in a), f'Unsupported {kind} attributes: {list(a)}')
        if kind == 'path' and not filtered_parent:
            commands = set(re.findall(r'[A-DF-Za-df-z]', e.get('d', '')))
            require(commands <= set('MmLlHhVvCcSsZz'), 'Only M/L/H/V/C/S/Z paths are supported')
            require(len(re.findall('[Mm]', e.get('d', ''))) <= 1, 'Compound paths need fill-rule review; not supported')
            require('d' in a, 'Path is missing d')
        if kind == 'rect':
            require(float(e.get('x', 0)) == 0 and float(e.get('y', 0)) == 0 and
                    float(e.get('width', 0)) == box[2] and float(e.get('height', 0)) == box[3] and
                    e.getparent() is root and not e.get('rx') and not e.get('ry'),
                    'Only a full-page background rect is supported; convert other rectangles to paths')
        if kind == 'linearGradient':
            require(e.get('gradientUnits') == 'userSpaceOnUse' and not e.get('gradientTransform') and
                    not any(ET.QName(k).localname == 'href' for k in e.attrib) and e.get('spreadMethod', 'pad') == 'pad',
                    'Only direct userSpaceOnUse linear gradients with pad are supported')
            points = [float(e.get(k)) for k in ('x1', 'y1', 'x2', 'y2')]
            require(points[:2] != points[2:], 'Zero-length gradient')
            offsets = [float(s.get('offset')) for s in e]
            require(len(offsets) >= 2 and offsets == sorted(offsets) and offsets[0] == 0 and offsets[-1] == 1,
                    'Gradient offsets must be ordered numeric values from 0 to 1')
            for stop in e:
                require(set(attrs(stop)) <= {'offset', 'stop-color'}, 'Gradient opacity/other stop styles are not supported')
        if kind == 'filter':
            require(e.get('filterUnits') == 'userSpaceOnUse' and len(e) == 1 and
                    ET.QName(e[0]).localname == 'feGaussianBlur' and
                    set(e[0].attrib) <= {'stdDeviation'}, 'Only bounded Gaussian-blur shadows are supported')
            require(all(k in e.attrib for k in ('x', 'y', 'width', 'height')), 'Shadow filter bounds required')
    return box


def sub(parent, tag, **attributes):
    e = OxmlElement(tag)
    for key, value in attributes.items():
        e.set(key, str(value))
    parent.append(e)
    return e


def color(parent, value):
    rgb = ImageColor.getrgb(value)
    require(len(rgb) == 3, 'RGBA paint is not supported')
    return sub(parent, 'a:srgbClr', val=''.join(f'{v:02X}' for v in rgb))


def extract_shadows(pdf, trace, groups, definitions, scale, work, run):
    """Pair rendered image/SMask in paint order; check dimensions, bounds and alpha."""
    images = trace.findall('.//fill_image')
    require(len(images) == len(groups), 'PDF image count differs from SVG shadows; inspect before continuing')
    if not groups:
        return []
    listing = run(['pdfimages', '-list', pdf], work)
    rows = [line.split() for line in listing.splitlines() if re.match(r'\s*\d+\s+\d+\s+', line)]
    require(len(rows) == len(groups)*2, 'Expected one image + one SMask per shadow')
    prefix = work / 'shadow'
    run(['pdfimages', '-png', pdf, prefix], work)
    result = []
    for i, (image, group) in enumerate(zip(images, groups)):
        rgbrow, maskrow = rows[i*2:i*2+2]
        require(rgbrow[2] == 'image' and maskrow[2] == 'smask' and rgbrow[3:5] == maskrow[3:5],
                'Unexpected PDF image/mask pair')
        require(rgbrow[10:12] == maskrow[10:12], 'SMask belongs to a different PDF image')
        w, h = int(image.get('width')), int(image.get('height'))
        require([w, h] == list(map(int, rgbrow[3:5])), 'PDF image paint/resource order mismatch')
        a,b,c,d,x,y = map(float, image.get('transform').split())
        require(abs(b)+abs(c) < 1e-6 and a > 0 and d > 0, 'Rotated/flipped shadow image not supported')
        filt = definitions[group.get('filter')[5:-1]]
        expected = [float(filt.get(k))*scale for k in ('x','y','width','height')]
        actual = [v*12700 for v in (x,y,a,d)]  # PDF points → EMU
        require(all(abs(u-v) < 2*9525 for u,v in zip(expected, actual)),
                'Shadow placement differs from filter bounds by more than 2 CSS pixels')
        alpha = float(image.get('alpha', 1)) * prod(float(p.get('alpha', 1)) for p in image.iterancestors())
        require(isclose(alpha, float(group.get('opacity', 1)), abs_tol=.005), 'Shadow opacity mismatch')
        rgb = Image.open(f'{prefix}-{int(rgbrow[1]):03}.png').convert('RGB')
        mask = Image.open(f'{prefix}-{int(maskrow[1]):03}.png').convert('L')
        require(rgb.size == mask.size == (w,h), 'Shadow image dimensions differ')
        rgb.putalpha(mask.point([round(v*alpha) for v in range(256)]))
        stream = BytesIO()
        rgb.save(stream, format='PNG')
        result.append((stream.getvalue(), tuple(round(v) for v in actual)))
    return result


def build(root, template, pdf, trace, target, work, run):
    prs = Presentation(template)
    require(len(prs.slides) == 1, 'Text template is empty or has multiple slides')
    slide = prs.slides[0]
    if prs.slide_width is None or prs.slide_height is None:
        raise ValueError('Text template lacks page dimensions')
    width, height = prs.slide_width/914400, prs.slide_height/914400
    require(all(isclose(a, b, rel_tol=1e-4) for a, b in zip((width, height), svg_size(root))),
            'Imported slide size differs from the source SVG')
    box = validate_svg(root, width, height)
    scale = prs.slide_width/box[2]
    texts = [deepcopy(s._element) for s in slide.shapes if isinstance(s, Shape) and s.text]
    definitions = {e.get('id'): e for e in root.iter() if e.get('id')}
    groups = [e for e in root.iter() if e.get('filter')]
    shadows = extract_shadows(pdf, trace, groups, definitions, scale, work, run)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    slide.shapes.turbo_add_enabled = True
    counts = Counter()

    def fill(sp, value):
        if value == 'none':
            sub(sp, 'a:noFill')
        elif value.startswith('url('):
            gradient = definitions[value[5:-1]]
            gf = sub(sp, 'a:gradFill', rotWithShape='1')
            stops = sub(gf, 'a:gsLst')
            for stop in gradient:
                gs = sub(stops, 'a:gs', pos=round(float(stop.get('offset'))*100000))
                color(gs, attrs(stop)['stop-color'])
            angle = degrees(atan2(float(gradient.get('y2'))-float(gradient.get('y1')),
                                  float(gradient.get('x2'))-float(gradient.get('x1')))) % 360
            # ponytail: preserve angle/stops; SVG gradient endpoint offsets may differ.
            # Add projected endpoint remapping only if visual review shows a material mismatch.
            sub(gf, 'a:lin', ang=round(angle*60000), scaled='0')
        else:
            color(sub(sp, 'a:solidFill'), value)

    def draw_path(e, label, paint):
        pen = RecordingPen()
        parse_path(e.get('d'), pen)
        bounds = BoundsPen(None)
        pen.replay(bounds)
        require(bounds.bounds is not None, 'Empty path; remove it explicitly before conversion')
        x0,y0,x1,y1 = bounds.bounds
        w,h = max(x1-x0, 1/scale), max(y1-y0, 1/scale)
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, round(x0*scale), round(y0*scale), round(w*scale), round(h*scale))
        sh.name = label or f'SVG path {counts["paths"]+1}'
        sp = sh._element.spPr
        for child in list(sp):
            if ET.QName(child).localname != 'xfrm':
                sp.remove(child)
        geom = sub(sp, 'a:custGeom')
        for tag in ('avLst','gdLst','ahLst','cxnLst'):
            sub(geom, 'a:'+tag)
        sub(geom, 'a:rect', l='0', t='0', r='r', b='b')
        path = sub(sub(geom, 'a:pathLst'), 'a:path', w=round(w*scale), h=round(h*scale))
        for op, points in pen.value:
            tag = {'moveTo':'moveTo', 'lineTo':'lnTo', 'curveTo':'cubicBezTo', 'closePath':'close', 'endPath':None}[op]
            if tag:
                node = sub(path, 'a:'+tag)
                for x,y in points:
                    sub(node, 'a:pt', x=round((x-x0)*scale), y=round((y-y0)*scale))
        fill(sp, paint['fill'])
        ln = sub(sp, 'a:ln', w=round(float(paint['stroke-width'])*scale),
                 cap={'round':'rnd', 'square':'sq', 'butt':'flat'}[paint['stroke-linecap']])
        fill(ln, paint['stroke'])
        join = paint['stroke-linejoin']
        sub(ln, 'a:'+{'round':'round','bevel':'bevel','miter':'miter'}[join], **({'lim':'400000'} if join=='miter' else {}))
        sub(sp, 'a:effectLst')
        for child in list(sh._element):
            if ET.QName(child).localname == 'style':
                sh._element.remove(child)
        counts['paths'] += 1

    def walk(e, label='', inherited=None):
        kind = ET.QName(e).localname
        if kind in ('defs','text','title','desc'):
            return
        paint = dict(inherited or {'fill':'black', 'stroke':'none', 'stroke-width':'1',
                                   'stroke-linecap':'butt', 'stroke-linejoin':'miter'})
        paint.update({k:v for k,v in attrs(e).items() if k in PAINT})
        label = e.get('data-name', label)
        if e.get('filter'):
            data, position = shadows[counts['shadows']]
            pic = slide.shapes.add_picture(BytesIO(data), *position)
            pic.name = f'Soft shadow {counts["shadows"]+1}'
            counts['shadows'] += 1
        elif kind == 'path':
            draw_path(e, label, paint)
        elif kind == 'rect':
            require(paint['stroke'] == 'none', 'Background rect stroke is unsupported')
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(*ImageColor.getrgb(paint['fill']))
        else:
            for child in e:
                walk(child, label, paint)

    walk(root)
    # This workflow intentionally keeps PDF-imported text above the geometry.
    # Occluding/overlapping artwork needs explicit z-order reconstruction instead.
    for element in texts:
        for node in element.iter():
            if node.get('typeface') == 'TimesNewRoman':
                node.set('typeface', 'Times New Roman')
        slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    for i,node in enumerate(slide.shapes._spTree.findall('.//{*}cNvPr'), 1):
        node.set('id', str(i))
    original = ''.join(root.xpath('.//*[local-name()="text"]//text()'))
    converted = ''.join(s.text for s in slide.shapes if isinstance(s, Shape))
    require(Counter(re.sub(r'\s+', '', original)) == Counter(re.sub(r'\s+', '', converted)),
            'Text character inventory differs; inspect fonts, shaping or import before proceeding')
    expected = len(root.findall('.//{*}path')) - len(groups)
    require(counts['paths'] == expected and counts['shadows'] == len(groups), 'Geometry/shadow count mismatch')
    prs.save(target)
    check = Presentation(target)
    require(len(check.slides[0].shapes) == expected + len(groups) + len(texts), 'PPTX object count mismatch')
    ids = [e.get('id') for e in check.slides[0].shapes._spTree.findall('.//{*}cNvPr')]
    require(len(ids) == len(set(ids)), 'Duplicate PPTX object IDs')
    with ZipFile(target) as z:
        require(z.testzip() is None, 'Invalid PPTX ZIP')
    return {'pptx': str(target), 'native_paths': expected, 'text_boxes': len(texts), 'shadow_images': len(groups),
            'text_check': 'Whitespace-normalized character inventory; not a word-order/layout check',
            'limitations': ['All text is above geometry', 'Gradient endpoints may differ; angle/stops retained',
                            'Some text/indices split across boxes; imported baselines may shift slightly', 'No original PPT layers/groups/animations reconstructed']}
